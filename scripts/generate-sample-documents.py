# -*- coding: utf-8 -*-
"""Sinh bộ tài liệu hành chính mẫu (tổng hợp/synthetic) phục vụ kiểm thử pipeline
trích xuất chỉ tiêu: PDF có text layer, PDF scan (ảnh), DOCX và XLSX.

Chạy:  python scripts/generate-sample-documents.py
Kết quả ghi vào thư mục samples/.
"""
import io
import os

SAMPLES_DIR = os.path.join(os.path.dirname(__file__), '..', 'samples')

KE_HOACH_TEXT = [
    "ỦY BAN NHÂN DÂN PHƯỜNG LÁI THIÊU",
    "Số: 21/KH-UBND",
    "",
    "KẾ HOẠCH",
    "Phát triển kinh tế - xã hội phường Lái Thiêu năm 2026",
    "",
    "I. MỤC ĐÍCH, YÊU CẦU",
    "Cụ thể hóa các nhiệm vụ trọng tâm năm 2026, phân công rõ đơn vị chủ trì,",
    "tiến độ và trách nhiệm báo cáo kết quả thực hiện.",
    "",
    "II. CHỈ TIÊU CHỦ YẾU NĂM 2026",
    "1. Tỷ lệ hồ sơ thủ tục hành chính giải quyết đúng hạn đạt 98% trở lên.",
    "Đơn vị chủ trì: Trung tâm Phục vụ hành chính công. Báo cáo hàng quý.",
    "2. Tổng thu ngân sách nhà nước trên địa bàn đạt 3.450 tỷ đồng,",
    "hoàn thành trước ngày 31/12/2026. Đơn vị chủ trì: Phòng Kinh tế, Hạ tầng và Đô thị.",
    "3. Tỷ lệ người dân tham gia bảo hiểm y tế đạt tối thiểu 95,5%.",
    "Đơn vị chủ trì: Phòng Văn hóa - Xã hội, phối hợp: Trạm Y tế phường. Báo cáo hàng quý.",
    "4. Trồng mới 1.200 cây xanh đô thị trên các tuyến đường chính.",
    "Đơn vị chủ trì: Phòng Kinh tế, Hạ tầng và Đô thị. Báo cáo hàng quý.",
    "5. Tỷ lệ hộ nghèo theo chuẩn đa chiều giảm còn 0,8% vào cuối năm 2026.",
    "Đơn vị chủ trì: Phòng Văn hóa - Xã hội. Báo cáo hàng quý.",
    "6. Giải quyết việc làm mới cho 2.500 người trong năm 2026.",
    "Đơn vị chủ trì: Phòng Văn hóa - Xã hội. Báo cáo hàng tháng.",
    "7. Tỷ lệ hồ sơ dịch vụ công trực tuyến toàn trình đạt 70%.",
    "Đơn vị chủ trì: Trung tâm Phục vụ hành chính công. Báo cáo hàng tháng.",
    "8. Số vụ phạm pháp hình sự trên địa bàn không quá 45 vụ.",
    "Đơn vị chủ trì: Công an phường. Báo cáo hàng tháng.",
    "",
    "III. TỔ CHỨC THỰC HIỆN",
    "Các đơn vị được giao chủ trì chịu trách nhiệm cập nhật số liệu về Văn phòng",
    "HĐND và UBND phường trước ngày 05 của tháng liền kề sau kỳ báo cáo.",
    "",
    "TM. ỦY BAN NHÂN DÂN",
    "CHỦ TỊCH",
]

BAO_CAO_ROWS = [
    ("STT", "Chỉ tiêu", "Đơn vị tính", "Kế hoạch năm 2026", "Thực hiện 6 tháng", "Đơn vị chủ trì"),
    (1, "Tỷ lệ hồ sơ TTHC giải quyết đúng hạn", "%", 98, 96.8, "Trung tâm Phục vụ hành chính công"),
    (2, "Tổng thu ngân sách nhà nước", "tỷ đồng", 3450, 1977.7, "Phòng Kinh tế, Hạ tầng và Đô thị"),
    (3, "Tỷ lệ người dân tham gia BHYT", "%", 95.5, 95.2, "Phòng Văn hóa - Xã hội"),
    (4, "Trồng mới cây xanh đô thị", "cây", 1200, 640, "Phòng Kinh tế, Hạ tầng và Đô thị"),
    (5, "Tỷ lệ hộ nghèo đa chiều", "%", 0.8, 1.1, "Phòng Văn hóa - Xã hội"),
    (6, "Giải quyết việc làm mới", "người", 2500, 1310, "Phòng Văn hóa - Xã hội"),
]

QUYET_DINH_TEXT = [
    "ỦY BAN NHÂN DÂN PHƯỜNG LÁI THIÊU",
    "Số: 145/QĐ-UBND",
    "",
    "QUYẾT ĐỊNH",
    "Về việc giao chỉ tiêu chuyển đổi số năm 2026",
    "",
    "Điều 1. Giao các chỉ tiêu chuyển đổi số năm 2026 như sau:",
    "1. Tỷ lệ số hóa hồ sơ, kết quả giải quyết thủ tục hành chính đạt 100%.",
    "Đơn vị chủ trì: Trung tâm Phục vụ hành chính công. Báo cáo hàng tháng.",
    "2. Tỷ lệ người dân trưởng thành có tài khoản định danh điện tử mức độ 2",
    "đạt tối thiểu 90%. Đơn vị chủ trì: Công an phường. Báo cáo hàng quý.",
    "3. Lắp đặt mới 25 điểm camera an ninh kết nối về trung tâm điều hành,",
    "hoàn thành trước ngày 30/09/2026. Đơn vị chủ trì: Công an phường.",
    "",
    "Điều 2. Quyết định này có hiệu lực kể từ ngày ký.",
]


def ensure_dir():
    os.makedirs(SAMPLES_DIR, exist_ok=True)


def make_docx():
    from docx import Document
    doc = Document()
    for line in KE_HOACH_TEXT:
        doc.add_paragraph(line)
    path = os.path.join(SAMPLES_DIR, 'ke-hoach-ktxh-2026.docx')
    doc.save(path)
    print('DOCX :', path)


def make_xlsx():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'CHI_TIEU_2026'
    ws.append(["BÁO CÁO KẾT QUẢ THỰC HIỆN CHỈ TIÊU 6 THÁNG ĐẦU NĂM 2026"])
    ws.append([])
    for row in BAO_CAO_ROWS:
        ws.append(list(row))
    path = os.path.join(SAMPLES_DIR, 'bao-cao-chi-tieu-6-thang-2026.xlsx')
    wb.save(path)
    print('XLSX :', path)


FONT_CANDIDATES = [
    r'C:\Windows\Fonts\times.ttf',
    r'C:\Windows\Fonts\arial.ttf',
    r'C:\Windows\Fonts\segoeui.ttf',
]


def find_font():
    for path in FONT_CANDIDATES:
        if os.path.exists(path):
            return path
    raise RuntimeError('Không tìm thấy font hỗ trợ tiếng Việt')


def make_pdf_text():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    font_path = find_font()
    pdfmetrics.registerFont(TTFont('VNFont', font_path))
    path = os.path.join(SAMPLES_DIR, 'quyet-dinh-chuyen-doi-so-2026.pdf')
    pdf = canvas.Canvas(path, pagesize=A4)
    width, height = A4
    y = height - 60
    pdf.setFont('VNFont', 12)
    for line in QUYET_DINH_TEXT:
        if y < 60:
            pdf.showPage()
            pdf.setFont('VNFont', 12)
            y = height - 60
        pdf.drawString(55, y, line)
        y -= 20
    pdf.save()
    print('PDF  :', path)


def make_pdf_scan():
    """PDF dạng scan: chữ được vẽ thành ảnh rồi nhúng vào PDF, không có text layer."""
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    font = ImageFont.truetype(find_font(), 34)
    image = Image.new('RGB', (1654, 2339), 'white')  # A4 ở 200 DPI
    draw = ImageDraw.Draw(image)
    y = 90
    for line in KE_HOACH_TEXT:
        draw.text((110, y), line, fill=(20, 20, 20), font=font)
        y += 52
    # Nhiễu nhẹ mô phỏng bản scan.
    image = image.rotate(-0.4, fillcolor='white', expand=False)
    path = os.path.join(SAMPLES_DIR, 'ke-hoach-ktxh-2026-scan.pdf')
    pdf = canvas.Canvas(path, pagesize=A4)
    buffer = io.BytesIO()
    image.save(buffer, format='JPEG', quality=82)
    buffer.seek(0)
    pdf.drawImage(ImageReader(buffer), 0, 0, width=A4[0], height=A4[1])
    pdf.save()
    print('SCAN :', path)


def make_png():
    from PIL import Image, ImageDraw, ImageFont
    font = ImageFont.truetype(find_font(), 30)
    lines = QUYET_DINH_TEXT
    image = Image.new('RGB', (1500, 120 + len(lines) * 48), 'white')
    draw = ImageDraw.Draw(image)
    y = 60
    for line in lines:
        draw.text((90, y), line, fill=(15, 15, 15), font=font)
        y += 48
    path = os.path.join(SAMPLES_DIR, 'quyet-dinh-chuyen-doi-so-2026-anh.png')
    image.save(path)
    print('PNG  :', path)


if __name__ == '__main__':
    ensure_dir()
    make_docx()
    make_xlsx()
    make_pdf_text()
    make_pdf_scan()
    make_png()
    print('Hoàn tất sinh tài liệu mẫu.')
