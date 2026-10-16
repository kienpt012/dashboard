"""Dựng tài liệu Hướng dẫn sử dụng IOC (.docx) từ các tệp Markdown trong thư mục noi-dung/.

Cách dùng:
    python docs/huong-dan-su-dung/build.py            # tạo HDSD-IOC.docx
    powershell -File docs/huong-dan-su-dung/xuat-pdf.ps1   # cập nhật mục lục và xuất HDSD-IOC.pdf bằng Word

Cú pháp Markdown được hỗ trợ (đủ cho tài liệu hướng dẫn):
    # Phần ...            đầu phần (sang trang mới)
    ## Mục ...            mục (vào mục lục)
    ### Tiểu mục          tiểu mục (vào mục lục)
    #### Nhãn             nhãn nhỏ trong mục (không vào mục lục)
    **đậm**, *nghiêng*, `Nút hoặc nhãn trên màn hình`
    - gạch đầu dòng (thụt 2 dấu cách cho cấp 2)
    1. bước thao tác (đánh số lại từ 1 cho mỗi danh sách)
    ![Chú thích hình](anh/ten-anh.png)       ảnh rộng hết khổ, kèm "Hình n. Chú thích hình"
    ![Chú thích hình|60](anh/ten-anh.png)    ảnh rộng 60% khổ
    > Ghi chú                                 khung ghi chú (xanh)
    > [!CANH-BAO] Nội dung                    khung cảnh báo (vàng cam)
    | Cột | Cột |  bảng, dòng đầu là tiêu đề, dòng thứ hai là |---|---|
    [[sang-trang]]                            ngắt trang
"""

from __future__ import annotations

import datetime as dt
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent
CONTENT = ROOT / "noi-dung"
OUTPUT = ROOT / "HDSD-IOC.docx"

FONT = "Segoe UI"
BRAND = RGBColor(0x1D, 0x4E, 0xD8)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
NOTE_FILL = "EEF4FF"
NOTE_LINE = "1D4ED8"
WARN_FILL = "FFF7E6"
WARN_LINE = "D97706"
TABLE_HEAD = "E8EEFC"
TABLE_LINE = "CBD5E1"

PAGE_WIDTH = Cm(21.0)
MARGIN_SIDE = Cm(2.0)
CONTENT_WIDTH = PAGE_WIDTH - 2 * MARGIN_SIDE

TITLE = "HƯỚNG DẪN SỬ DỤNG"
SUBTITLE = "Hệ thống Trung tâm điều hành thông minh (IOC) cấp xã"
TAGLINE = "Từ dữ liệu đến trách nhiệm"


# --------------------------------------------------------------------------- tiện ích XML

def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, **edges: tuple[str, int]) -> None:
    """edges: top/left/bottom/right=(màu, độ dày theo 1/8 pt); màu 'nil' để bỏ viền."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge, (color, size) in edges.items():
        element = OxmlElement(f"w:{edge}")
        if color == "nil":
            element.set(qn("w:val"), "nil")
        else:
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), str(size))
            element.set(qn("w:color"), color)
        borders.append(element)
    tc_pr.append(borders)


def set_cell_margins(cell, top=80, bottom=80, left=140, right=140) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tc_pr.append(margins)


def set_table_borders(table, color: str = TABLE_LINE, size: int = 4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)
        borders.append(element)
    tbl_pr.append(borders)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:cantSplit")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run._r.append(begin)
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run._r.append(instr)
    run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)
    paragraph.add_run(placeholder)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def paragraph_border(paragraph, edge: str, color: str, size: int = 6, space: int = 4) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), str(space))
    element.set(qn("w:color"), color)
    borders.append(element)


# --------------------------------------------------------------------------- kiểu chữ

def setup_styles(document: Document) -> None:
    styles = document.styles

    def font(style, size: float, bold=False, color=INK, italic=False):
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = color
        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.append(r_fonts)
        for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            r_fonts.set(qn(key), FONT)

    normal = styles["Normal"]
    font(normal, 10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, (size, before, after) in {1: (22, 0, 14), 2: (15, 18, 8), 3: (12, 14, 6), 4: (10.5, 10, 4)}.items():
        style = styles[f"Heading {level}"]
        font(style, size, bold=True, color=BRAND if level < 4 else INK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.1
    styles["Heading 1"].paragraph_format.page_break_before = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[name]
        font(style, 10.5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.25

    caption = styles["Caption"]
    font(caption, 9, italic=True, color=MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(12)

    for level in (1, 2, 3):
        name = f"TOC {level}"
        try:
            style = styles[name]
        except KeyError:
            from docx.enum.style import WD_STYLE_TYPE
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = normal
        font(style, 10.5 if level == 1 else 10, bold=level == 1)
        style.paragraph_format.space_before = Pt(6 if level == 1 else 0)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.left_indent = Cm(0.6 * (level - 1))


# --------------------------------------------------------------------------- nội dung trong dòng

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")


def add_inline(paragraph, text: str, base_bold=False, color: RGBColor | None = None) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            # Nhãn, nút trên màn hình: đậm, màu thương hiệu
            run = paragraph.add_run(part[1:-1])
            run.bold = True
            run.font.color.rgb = BRAND
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
            if base_bold:
                run.bold = True
        if color is not None and not (part.startswith("`") and part.endswith("`")):
            run.font.color.rgb = color


# --------------------------------------------------------------------------- khối

CHAR_WIDTH = Cm(0.205)  # bề rộng trung bình một ký tự Segoe UI 10,5pt (chữ đậm ở dòng tiêu đề)
CELL_PADDING = Cm(0.4)


def column_widths(rows: list[list[str]], columns: int) -> list[Emu]:
    """Cột có chữ ngắn (≤ 18 ký tự, hoặc từ dài nhất) được đủ chỗ trên một dòng; phần còn lại chia theo độ dài chữ."""
    def plain(text: str) -> str:
        return re.sub(r"[`*]", "", text)

    # Tiêu đề được phép xuống dòng: chỉ tính từ dài nhất của tiêu đề, độ dài cả ô tính theo thân bảng.
    cells = [[plain(row[i]) if i < len(row) else "" for row in rows[1:]] for i in range(columns)]
    heads = [plain(rows[0][i]) if i < len(rows[0]) else "" for i in range(columns)]
    longest = [max([len(text) for text in column] + [max((len(word) for word in heads[i].split()), default=0)]) for i, column in enumerate(cells)]
    longest_word = [max((len(word) for text in column + [heads[i]] for word in text.split()), default=0) for i, column in enumerate(cells)]
    minimum = []
    for index in range(columns):
        chars = longest[index] if longest[index] <= 18 else max(longest_word[index], 10)
        minimum.append(int(CHAR_WIDTH * chars + CELL_PADDING))
    total = int(CONTENT_WIDTH)
    if sum(minimum) >= total:
        return [Emu(int(total * value / sum(minimum))) for value in minimum]
    desire = [max(0, min(longest[index], 80) * int(CHAR_WIDTH) + int(CELL_PADDING) - minimum[index]) for index in range(columns)]
    spare = total - sum(minimum)
    if sum(desire) == 0:
        return [Emu(value + spare // columns) for value in minimum]
    return [Emu(minimum[index] + int(spare * desire[index] / sum(desire))) for index in range(columns)]

class Builder:
    def __init__(self, document: Document):
        self.doc = document
        self.figure = 0
        self.missing: list[str] = []
        self.list_num_id: int | None = None
        self.number_abstract = self._abstract_num_for("List Number")

    # Mỗi danh sách số bắt đầu lại từ 1: tạo một w:num mới trỏ tới cùng abstractNum, đặt lại số bắt đầu.
    def _abstract_num_for(self, style_name: str) -> str | None:
        style = self.doc.styles[style_name]
        num_pr = style.element.pPr.find(qn("w:numPr")) if style.element.pPr is not None else None
        if num_pr is None:
            return None
        num_id = num_pr.find(qn("w:numId")).get(qn("w:val"))
        numbering = self.doc.part.numbering_part.element
        for num in numbering.findall(qn("w:num")):
            if num.get(qn("w:numId")) == num_id:
                return num.find(qn("w:abstractNumId")).get(qn("w:val"))
        return None

    def _new_number_list(self) -> int:
        numbering = self.doc.part.numbering_part.element
        existing = [int(num.get(qn("w:numId"))) for num in numbering.findall(qn("w:num"))]
        num_id = max(existing + [0]) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract = OxmlElement("w:abstractNumId")
        abstract.set(qn("w:val"), self.number_abstract or "0")
        num.append(abstract)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "1")
        override.append(start)
        num.append(override)
        numbering.append(num)
        return num_id

    def heading(self, level: int, text: str) -> None:
        paragraph = self.doc.add_heading(level=level)
        add_inline(paragraph, text)
        if level == 1:
            paragraph_border(paragraph, "bottom", NOTE_LINE, size=12, space=8)

    def paragraph(self, text: str) -> None:
        paragraph = self.doc.add_paragraph()
        add_inline(paragraph, text)

    def bullet(self, text: str, level: int) -> None:
        paragraph = self.doc.add_paragraph(style="List Bullet 2" if level > 0 else "List Bullet")
        add_inline(paragraph, text)

    def numbered(self, text: str, fresh: bool) -> None:
        if fresh or self.list_num_id is None:
            self.list_num_id = self._new_number_list()
        paragraph = self.doc.add_paragraph(style="List Number")
        num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = self.list_num_id
        add_inline(paragraph, text)

    def image(self, caption: str, path: str, percent: int | None) -> None:
        file = (ROOT / path).resolve()
        if not file.exists():
            # Ảnh chưa chụp: để khung chờ cho bản nháp, báo ở cuối lần dựng.
            self.missing.append(path)
            self.figure += 1
            table = self.doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.rows[0].cells[0]
            cell.width = CONTENT_WIDTH
            set_cell_fill(cell, "F1F5F9")
            set_cell_margins(cell, top=600, bottom=600)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(cell.paragraphs[0], f"[Chưa có ảnh: {path}]", color=MUTED)
            cap = self.doc.add_paragraph(style="Caption")
            cap.add_run(f"Hình {self.figure}. ").bold = True
            add_inline(cap, caption)
            return
        with Image.open(file) as img:
            width_px, height_px = img.size
        width = Emu(int(CONTENT_WIDTH * (percent or 100) / 100))
        # Ảnh rất cao (chụp cả trang) không vượt quá 70% chiều cao vùng in.
        max_height = Cm(29.7 - 2.2 - 2.2) * 0.7
        if width * height_px / width_px > max_height:
            width = Emu(int(max_height * width_px / height_px))
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run()
        picture = run.add_picture(str(file), width=width)
        # Viền mảnh quanh ảnh chụp màn hình
        sp_pr = picture._inline.graphic.graphicData.pic.spPr
        line = OxmlElement("a:ln")
        line.set("w", "6350")
        fill = OxmlElement("a:solidFill")
        color = OxmlElement("a:srgbClr")
        color.set("val", "CBD5E1")
        fill.append(color)
        line.append(fill)
        sp_pr.append(line)
        self.figure += 1
        cap = self.doc.add_paragraph(style="Caption")
        cap.add_run(f"Hình {self.figure}. ").bold = True
        add_inline(cap, caption)

    def callout(self, lines: list[str]) -> None:
        warn = bool(lines) and lines[0].startswith("[!CANH-BAO]")
        if warn:
            lines = [lines[0][len("[!CANH-BAO]"):].strip(), *lines[1:]]
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        cell.width = CONTENT_WIDTH
        set_cell_fill(cell, WARN_FILL if warn else NOTE_FILL)
        line = WARN_LINE if warn else NOTE_LINE
        set_cell_borders(cell, left=(line, 24), top=("nil", 0), bottom=("nil", 0), right=("nil", 0))
        set_cell_margins(cell, top=100, bottom=100, left=200, right=160)
        first = True
        for text in lines:
            paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            if text.startswith("- "):
                paragraph.paragraph_format.left_indent = Cm(0.4)
                paragraph.paragraph_format.first_line_indent = Cm(-0.3)
                text = "• " + text[2:]
            add_inline(paragraph, text)
            first = False
        keep_row_together(table.rows[0])
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        spacer.paragraph_format.line_spacing = 0.6

    def table(self, rows: list[list[str]]) -> None:
        header, body = rows[0], rows[1:]
        columns = len(header)
        table = self.doc.add_table(rows=1, cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)
        widths = column_widths(rows, columns)
        for index, text in enumerate(header):
            cell = table.rows[0].cells[index]
            cell.width = widths[index]
            set_cell_fill(cell, TABLE_HEAD)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, text, base_bold=True)
        set_repeat_header(table.rows[0])
        for row in body:
            cells = table.add_row().cells
            for index in range(columns):
                cell = cells[index]
                cell.width = widths[index]
                set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.15
                value = row[index] if index < len(row) else ""
                if value in ("✓", "—", "-", "x", "X"):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, value)
            keep_row_together(table.rows[-1])
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# --------------------------------------------------------------------------- đọc Markdown

IMAGE = re.compile(r"^!\[(?P<caption>[^\]|]*)(?:\|(?P<percent>\d+))?\]\((?P<path>[^)]+)\)\s*$")


LIST_MARKER = re.compile(r"^\s*([-*]|\d+[.)])\s+")


def continued(lines: list[str], index: int, text: str) -> tuple[str, int]:
    """Gộp các dòng thụt đầu dòng ngay sau một mục danh sách vào mục đó; trả về nội dung và chỉ số dòng kế tiếp."""
    parts = [text.strip()]
    index += 1
    while index < len(lines):
        line = lines[index]
        if not line.strip() or not line[:1].isspace() or LIST_MARKER.match(line):
            break
        parts.append(line.strip())
        index += 1
    return " ".join(parts), index


def render_markdown(builder: Builder, text: str) -> None:
    lines = text.splitlines()
    index = 0
    paragraph: list[str] = []
    previous_kind = None

    def flush():
        nonlocal paragraph
        if paragraph:
            builder.paragraph(" ".join(part.strip() for part in paragraph))
            paragraph = []

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped or stripped.startswith("<!--"):
            flush()
            if stripped.startswith("<!--"):
                while index < len(lines) and "-->" not in lines[index]:
                    index += 1
            index += 1
            if not stripped:
                previous_kind = "blank" if previous_kind != "number" else "number-gap"
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            flush()
            builder.heading(len(heading.group(1)), heading.group(2))
            previous_kind = "heading"
        elif stripped == "[[sang-trang]]":
            flush()
            builder.page_break()
            previous_kind = "break"
        elif IMAGE.match(stripped):
            flush()
            match = IMAGE.match(stripped)
            builder.image(match.group("caption").strip(), match.group("path").strip(), int(match.group("percent")) if match.group("percent") else None)
            previous_kind = "image"
        elif stripped.startswith(">"):
            flush()
            block: list[str] = []
            joinable = False
            while index < len(lines) and lines[index].strip().startswith(">"):
                item = lines[index].strip()[1:].strip()
                if not item:
                    joinable = False
                elif item.startswith("- ") or not joinable:
                    block.append(item)
                    joinable = True
                else:
                    block[-1] += " " + item
                index += 1
            builder.callout(block)
            previous_kind = "callout"
            continue
        elif stripped.startswith("|"):
            flush()
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = lines[index].strip().strip("|")
                cells = [cell.strip() for cell in re.split(r"(?<!\\)\|", row)]
                if not all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells):
                    rows.append([cell.replace("\\|", "|") for cell in cells])
                index += 1
            builder.table(rows)
            previous_kind = "table"
            continue
        elif re.match(r"^\s*[-*]\s+", line):
            flush()
            level = 1 if len(line) - len(line.lstrip()) >= 2 else 0
            text, index = continued(lines, index, re.sub(r"^\s*[-*]\s+", "", line))
            builder.bullet(text, level)
            previous_kind = "bullet"
            continue
        elif re.match(r"^\d+[.)]\s+", stripped):
            flush()
            fresh = previous_kind not in ("number", "number-gap")
            text, index = continued(lines, index, re.sub(r"^\d+[.)]\s+", "", stripped))
            builder.numbered(text, fresh)
            previous_kind = "number"
            continue
        else:
            paragraph.append(stripped)
            previous_kind = "text"
        index += 1
    flush()


# --------------------------------------------------------------------------- bìa, mục lục, đầu trang

def cover(document: Document, version: str) -> None:
    section = document.sections[0]
    for _ in range(5):
        document.add_paragraph()
    label = document.add_paragraph()
    run = label.add_run(TAGLINE.upper())
    run.font.size = Pt(10)
    run.font.color.rgb = BRAND
    run.bold = True
    run.font.spacing = Pt(2)
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run(TITLE)
    run.font.size = Pt(34)
    run.bold = True
    run.font.color.rgb = INK
    subtitle = document.add_paragraph()
    paragraph_border(subtitle, "bottom", NOTE_LINE, size=18, space=12)
    subtitle.paragraph_format.space_after = Pt(24)
    run = subtitle.add_run(SUBTITLE)
    run.font.size = Pt(16)
    run.font.color.rgb = BRAND
    hero = next((path for path in (ROOT / "anh" / "bia.jpg", ROOT / "anh" / "bia.png") if path.exists()), None)
    if hero:
        picture = document.add_paragraph()
        picture.add_run().add_picture(str(hero), width=CONTENT_WIDTH)
    for text, size, color in (
        (f"Phiên bản {version}", 11, INK),
        (f"Cập nhật {dt.date.today().strftime('%d/%m/%Y')}", 10, MUTED),
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(text)
        run.font.size = Pt(size)
        run.font.color.rgb = color
    section.different_first_page_header_footer = True


def table_of_contents(document: Document) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.space_after = Pt(14)
    paragraph_border(heading, "bottom", NOTE_LINE, size=12, space=8)
    run = heading.add_run("MỤC LỤC")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = BRAND
    paragraph = document.add_paragraph()
    add_field(paragraph, 'TOC \\o "1-3" \\h \\z \\u', "Mở tệp bằng Word và chọn Cập nhật mục lục (hoặc chạy xuat-pdf.ps1).")


def header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(f"IOC  ·  {TITLE.capitalize()}")
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    paragraph_border(header, "bottom", TABLE_LINE, size=4, space=4)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for part in ("Trang ", None, " / ", "NUMPAGES"):
        if part is None:
            add_field(footer, "PAGE", "1")
        elif part == "NUMPAGES":
            add_field(footer, "NUMPAGES", "1")
        else:
            run = footer.add_run(part)
    for run in footer.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    args = [arg for arg in sys.argv[1:] if not arg.startswith("--")]
    version = args[0] if args else "2.0"
    document = Document()
    section = document.sections[0]
    section.page_width = PAGE_WIDTH
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = MARGIN_SIDE
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    setup_styles(document)
    core = document.core_properties
    core.title = f"{TITLE.capitalize()} — {SUBTITLE}"
    core.subject = TAGLINE
    core.language = "vi-VN"

    cover(document, version)
    header_footer(document)
    table_of_contents(document)

    builder = Builder(document)
    files = sorted(CONTENT.glob("*.md"))
    if not files:
        print("Không có tệp nội dung trong noi-dung/", file=sys.stderr)
        return 1
    for file in files:
        render_markdown(builder, file.read_text(encoding="utf-8"))
    document.save(OUTPUT)
    print(f"Đã tạo {OUTPUT.relative_to(ROOT.parent.parent)} — {len(files)} tệp nội dung, {builder.figure} hình")
    if builder.missing:
        print(f"Còn {len(builder.missing)} ảnh chưa có: " + ", ".join(builder.missing))
        return 2 if "--strict" in sys.argv else 0
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
